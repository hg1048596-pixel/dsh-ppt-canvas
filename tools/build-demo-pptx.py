"""Build the fixture .pptx used by dsh-ppt-canvas engine tests.

Reads tools/demo-deck.json and reproduces it with python-pptx as a REFERENCE
implementation only (tests / fixtures). The plugin itself never calls Python.

Usage:
  <python> tools/build-demo-pptx.py <out.pptx> [spec.json]
"""
from __future__ import annotations

import json
import sys
import os
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt
from PIL import Image, ImageDraw


def rgb(value: str | None):
    if not value:
        return None
    return RGBColor.from_string(value.upper().lstrip('#'))


def align_of(name: str | None):
    return {
        'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER,
        'right': PP_ALIGN.RIGHT, 'justify': PP_ALIGN.JUSTIFY,
    }.get(name or 'left', PP_ALIGN.LEFT)


def anchor_of(name: str | None):
    return {
        'top': MSO_ANCHOR.TOP, 'middle': MSO_ANCHOR.MIDDLE, 'bottom': MSO_ANCHOR.BOTTOM,
    }.get(name or 'top', MSO_ANCHOR.TOP)


def add_text(slide, spec):
    box = slide.shapes.add_textbox(Inches(spec['x']), Inches(spec['y']), Inches(spec['w']), Inches(spec['h']))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor_of(spec.get('anchor'))
    lines = spec['text'].split('\n')
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align_of(spec.get('align'))
        if spec.get('lineSpacing'):
            p.line_spacing = spec['lineSpacing']
        run = p.add_run()
        run.text = line
        f = run.font
        f.size = Pt(spec.get('size', 18))
        f.bold = bool(spec.get('bold'))
        f.italic = bool(spec.get('italic'))
        f.name = spec.get('font', 'Microsoft YaHei')
        c = rgb(spec.get('color', '000000'))
        if c is not None:
            f.color.rgb = c
    return box


def add_shape(slide, spec):
    kind = spec['type']
    if kind == 'rect':
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(spec['x']), Inches(spec['y']),
                                     Inches(spec['w']), Inches(spec['h']))
    elif kind == 'ellipse':
        shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(spec['x']), Inches(spec['y']),
                                     Inches(spec['w']), Inches(spec['h']))
    else:
        raise ValueError(f'unknown shape type {kind}')
    fill = rgb(spec.get('fill'))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    line = rgb(spec.get('line'))
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(spec.get('lineWidth', 1))
    shp.shadow.inherit = False
    return shp


def add_line(slide, spec):
    from pptx.enum.shapes import MSO_CONNECTOR
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(spec['x']), Inches(spec['y']),
                                      Inches(spec['x'] + spec['w']), Inches(spec['y'] + spec['h']))
    conn.line.color.rgb = rgb(spec.get('color', '000000'))
    conn.line.width = Pt(spec.get('lineWidth', 1))
    return conn


def add_table(slide, spec):
    rows = spec['rows']
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(spec['x']), Inches(spec['y']),
                                   Inches(spec['w']), Inches(spec['h']))
    table = shape.table
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(value)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(14)
                    run.font.name = 'Microsoft YaHei'
                    run.font.bold = r == 0
    return shape


def make_chart_png(path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    img = Image.new('RGB', (720, 480), (255, 255, 255))
    d = ImageDraw.Draw(img)
    bars = [('Q1', 120), ('Q2', 210), ('Q3', 390), ('Q4', 330)]
    base = 420
    for i, (label, h) in enumerate(bars):
        x = 80 + i * 160
        d.rectangle([x, base - h, x + 96, base], fill=(26, 115, 232))
        d.text((x + 30, base + 12), label, fill=(60, 64, 67))
    d.line([60, base, 680, base], fill=(150, 150, 150), width=2)
    d.text((70, 24), 'selection-instruction share (%)', fill=(32, 33, 36))
    img.save(path)


def build(spec_path: Path, out_path: Path) -> None:
    spec = json.loads(spec_path.read_text(encoding='utf-8'))
    prs = Presentation()
    prs.slide_width = Inches(spec['slideWidthIn'])
    prs.slide_height = Inches(spec['slideHeightIn'])
    blank = prs.slide_layouts[6]
    media_dir = out_path.parent / 'media'
    for slide_spec in spec['slides']:
        slide = prs.slides.add_slide(blank)
        bg = rgb(slide_spec.get('background', 'FFFFFF'))
        if bg is not None:
            from pptx.oxml.ns import qn
            from lxml import etree
            cSld = slide._element.find(qn('p:cSld'))
            bgPr = etree.SubElement(cSld, qn('p:bg'))
            bgPrPr = etree.SubElement(bgPr, qn('p:bgPr'))
            fill = etree.SubElement(bgPrPr, qn('a:solidFill'))
            clr = etree.SubElement(fill, qn('a:srgbClr'))
            clr.set('val', str(bg))
            etree.SubElement(bgPrPr, qn('a:effectLst'))
            cSld.insert(0, bgPr)
        for shape_spec in slide_spec['shapes']:
            kind = shape_spec['type']
            if kind == 'text':
                add_text(slide, shape_spec)
            elif kind in ('rect', 'ellipse'):
                add_shape(slide, shape_spec)
            elif kind == 'line':
                add_line(slide, shape_spec)
            elif kind == 'table':
                add_table(slide, shape_spec)
            elif kind == 'image':
                png = media_dir / shape_spec['image']
                make_chart_png(png)
                slide.shapes.add_picture(str(png), Inches(shape_spec['x']), Inches(shape_spec['y']),
                                         Inches(shape_spec['w']), Inches(shape_spec['h']))
            else:
                raise ValueError(f'unknown type {kind}')
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    print(f'wrote {out_path} ({out_path.stat().st_size} bytes)')


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print(__doc__)
        raise SystemExit(2)
    out = Path(sys.argv[1]).resolve()
    spec = Path(sys.argv[2]).resolve() if len(sys.argv) > 2 else Path(__file__).with_name('demo-deck.json')
    build(spec, out)
